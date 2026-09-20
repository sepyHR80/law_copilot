"""Generate test parsing fixtures."""

from pathlib import Path

FIXTURES = Path(__file__).parent / "fixtures"
FIXTURES.mkdir(exist_ok=True)

# Sample PDF
import pymupdf
doc = pymupdf.open()
page = doc.new_page()
page.insert_text((72, 72), "This is a sample PDF document.\nPage 1 content.")
page2 = doc.new_page()
page2.insert_text((72, 72), "Page 2 content.")
page3 = doc.new_page()
# Empty page
doc.save(str(FIXTURES / "sample.pdf"))
doc.close()

# Multi-page PDF
doc = pymupdf.open()
for i in range(5):
    page = doc.new_page()
    page.insert_text((72, 72), f"Multi-page PDF - Page {i + 1}\nLegal content here.")
doc.save(str(FIXTURES / "multi_page.pdf"))
doc.close()

# Empty page PDF
doc = pymupdf.open()
doc.new_page()  # empty
page = doc.new_page()
page.insert_text((72, 72), "Page 2 has content.")
doc.save(str(FIXTURES / "empty_page.pdf"))
doc.close()

# Invalid PDF
(FIXTURES / "invalid.pdf").write_text("This is not a valid PDF file.")

# Sample TXT
(FIXTURES / "sample.txt").write_text(
    "This is a sample text document.\n"
    "Article 1: General Provisions\n"
    "This is the complete text content.\n"
    "It preserves UTF-8 characters: \u06cc\u0627 \u0639\u0644\u064a\u0643\u0645\n",
    encoding="utf-8",
)

# Sample HTML
(FIXTURES / "sample.html").write_text(
    """<!DOCTYPE html>
<html>
<head><title>Test HTML Document</title></head>
<body>
<h1>Legal Document</h1>
<p>This is a paragraph of legal text.</p>
<script>alert('evil');</script>
<style>.hidden { display: none; }</style>
<h2>Section 1</h2>
<p>More content here.</p>
<ul><li>Item 1</li><li>Item 2</li></ul>
<table><tr><td>Cell A</td><td>Cell B</td></tr></table>
</body>
</html>""",
    encoding="utf-8",
)

# Sample DOCX
from docx import Document as DocxDocument

doc = DocxDocument()
doc.add_heading("Document Title", level=1)
doc.add_heading("Section 1", level=2)
doc.add_paragraph("This is a paragraph in a DOCX document.")
doc.add_paragraph("Another paragraph with legal content.")
table = doc.add_table(rows=2, cols=2)
table.rows[0].cells[0].text = "Header A"
table.rows[0].cells[1].text = "Header B"
table.rows[1].cells[0].text = "Value 1"
table.rows[1].cells[1].text = "Value 2"
doc.save(str(FIXTURES / "sample.docx"))

print("Fixtures created successfully.")
