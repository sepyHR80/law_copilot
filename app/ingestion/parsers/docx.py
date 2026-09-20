"""DOCX parser using python-docx."""

from io import BytesIO
from app.ingestion.parsers.base import DocumentParser
from app.ingestion.parsers.models import ParsedDocument, ParsedPage, ParsedTable
from app.ingestion.parsers.exceptions import DocumentParsingError


class DOCXParser(DocumentParser):
    """Parser for DOCX documents using python-docx."""

    def parse(self, content: bytes, document_id: str, document_version_id: str) -> ParsedDocument:
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(BytesIO(content))
        except Exception as exc:
            raise DocumentParsingError(f"Failed to open DOCX: {exc}") from exc

        # Extract all paragraphs and headings in order
        text_parts = []
        for para in doc.paragraphs:
            style = para.style.name if para.style else ""
            text = para.text.strip()
            if not text:
                continue
            if "Heading" in style:
                text_parts.append(f"[{style}] {text}")
            else:
                text_parts.append(text)

        # Extract tables
        tables = []
        for table in doc.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            tables.append(ParsedTable(rows=rows, metadata={}))

        full_text = "\n".join(text_parts)

        return ParsedDocument(
            document_id=document_id,
            document_version_id=document_version_id,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            pages=[ParsedPage(page_number=1, text=full_text, metadata={"type": "docx"})],
            tables=tables,
            metadata={"paragraph_count": len(doc.paragraphs), "table_count": len(tables)},
        )
